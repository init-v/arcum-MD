"""PPTX extractor — uses python-pptx directly, no ML models required."""

from __future__ import annotations

from pathlib import Path
from typing import Callable, Optional

from arcum.output import ConversionResult


def extract(
    file: Path,
    status_fn: Optional[Callable[[str], None]] = None,
) -> ConversionResult:
    """
    Convert a PPTX to Markdown using python-pptx directly.

    Produces structured Markdown from slide decks:
    - Slide titles become H2 headings
    - Bullet points become indented lists (level-aware)
    - Tables rendered as Markdown tables
    - Slides separated by horizontal rules
    """
    from pptx import Presentation

    def status(msg: str) -> None:
        if status_fn:
            status_fn(msg)

    size_mb = file.stat().st_size / 1_048_576
    status(f"Detected: {file.name} ({size_mb:.1f} MB)")
    status("Extracting slides...")

    prs = Presentation(str(file))
    slides = list(prs.slides)

    parts = [_slide_to_markdown(slide, i) for i, slide in enumerate(slides, start=1)]
    markdown = "\n\n---\n\n".join(p for p in parts if p.strip())

    slide_count = len(slides)
    status(f"Extracted {len(markdown.split()):,} words from {slide_count} slides.")

    return ConversionResult(
        source_path=file,
        file_type="pptx",
        body=markdown,
        metadata={"slides": slide_count},
    )


def _slide_to_markdown(slide, slide_num: int) -> str:
    parts: list[str] = []

    # Title
    title_shape = slide.shapes.title
    if title_shape and title_shape.has_text_frame:
        title_text = title_shape.text_frame.text.strip()
        parts.append(f"## {title_text}" if title_text else f"## Slide {slide_num}")
    else:
        parts.append(f"## Slide {slide_num}")

    # All other shapes in document order
    for shape in slide.shapes:
        if shape is title_shape:
            continue
        if shape.has_table:
            parts.append(_table_to_markdown(shape.table))
        elif shape.has_text_frame:
            block = _text_frame_to_markdown(shape.text_frame)
            if block:
                parts.append(block)

    return "\n\n".join(p for p in parts if p.strip())


def _text_frame_to_markdown(tf) -> str:
    lines: list[str] = []
    for para in tf.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        level = para.level
        if level == 0:
            lines.append(text)
        else:
            lines.append("  " * (level - 1) + f"- {text}")
    return "\n".join(lines)


def _table_to_markdown(table) -> str:
    rows: list[str] = []
    for i, row in enumerate(table.rows):
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
        if i == 0:
            rows.append("|" + "|".join([" --- "] * len(cells)) + "|")
    return "\n".join(rows)
